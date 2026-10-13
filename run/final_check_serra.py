import subprocess, re, hashlib, glob, os, zipfile
from docx import Document
R=[]; ok=lambda n,c,d="": R.append(("✅" if c else "❌",n,d))
sh=lambda c: subprocess.run(c,shell=True,capture_output=True,text=True).stdout
S="SERRA_submission/"; md5=lambda p: hashlib.md5(open(p,"rb").read()).hexdigest()
pdf=sh(f"pdftotext {S}Manuscript.pdf -"); pn=re.sub(r"\s+"," ",pdf)
D=Document(S+"Manuscript.docx"); dt=" ".join(p.text for p in D.paragraphs); dn=re.sub(r"\s+"," ",dt)
cl=" ".join(p.text for p in Document(S+"Cover_letter.docx").paragraphs)
TITLE="Decomposing environmental predictability into endogenous and exogenous components: theory, estimation and cross-domain evidence"
wt=[p.text for p in D.paragraphs if p.style.name.lower()=="title"]
ok("Title: PDF",TITLE in pn); ok("Title: Word",bool(wt) and wt[0]==TITLE); ok("Title: cover letter",TITLE in cl)
for nm,t in [("PDF",pn),("Word",dn),("Cover letter",cl)]:
    ok(f"{nm}: Qilu address + e-mail","2 Wenbo Road, Zhangqiu District, Jinan 250200" in t and "godstrength@qlnu.edu.cn" in t)
    ok(f"{nm}: no Kunsan e-mail","kunsan.ac.kr" not in t)
ok("Old EMS title absent (PDF/Word)","separating endogenous inertia from exogenous forcing" not in pn+dn)
a=pdf.split("Abstract",1)[1].split("Keywords",1)[0]; aw=len([w for w in re.sub(r"(?m)^\s*\d+\s*$","",a).split() if not w.isdigit()])
ok("Abstract 150-250 words",150<=aw<=250,f"{aw}")
kw=re.sub(r"\s+"," ",pdf.split("Keywords:",1)[1].split("Acknowledgements")[0]); k=[x for x in kw.split(";") if x.strip()]
ok("Keywords 4-6",4<=len(k)<=6,f"{len(k)}")
i=pdf.find("Introduction"); j=pdf.find("Related work",i); iw=len([w for w in re.sub(r"(?m)^\s*\d+\s*$","",pdf[i:j]).split() if not w.isdigit()])
ok("Introduction <= one page",iw<=550,f"{iw} words")
ok("Acknowledgements on title page","Acknowledgements" in sh(f"pdftotext -f 1 -l 2 {S}Manuscript.pdf -"))
ok("AI use documented in Methods","Use of AI tools" in pn and "Claude (Anthropic)" in pn)
ok("Statements and Declarations after References",pn.rfind("References")<pn.rfind("Statements and Declarations"))
for s in ["Funding","Competing interests","Author contributions","Data availability","Code availability"]: ok(f"Declaration: {s}",s in pn.split("Statements and Declarations")[-1])
ok("Unpublished B1 not in reference list","Persistence-anchored" not in pn.split("References")[-1].split("Statements and Declarations")[0])
ok("Highlights / graphical abstract absent","Highlights" not in pn and "Graphical abstract" not in pn)
eqp=sorted(set(int(x) for x in re.findall(r"(?m)^\((\d+)\)$",pdf))); ok("PDF: equations 1-24",eqp==list(range(1,25)),str(len(eqp)))
eqd=sorted(set(int(x) for x in re.findall(r"\\quad\\quad\\quad\\quad\((\d+)\)",sh(f"pandoc {S}Manuscript.docx -t latex 2>/dev/null")))); ok("Word: equations 1-24",eqd==list(range(1,25)),str(len(eqd)))
figs=sorted(set(int(x) for x in re.findall(r"Fig\. (\d+) ",pdf))); ok("PDF: Fig. 1-10 captions",set(range(1,11))<=set(figs))
ok("PDF: Tables 1-10 + A.1",all(f"Table {n}" in pn for n in range(1,11)) and "Table A.1" in pn)
ok("PDF: all figures vector",sh(f"pdfimages -list {S}Manuscript.pdf | tail -n +3 | wc -l").strip()=="0")
nb=len(re.findall(r"\\bibitem","".join(open("serra/Manuscript.bbl").read()))); nw=sum(1 for p in D.paragraphs if p.style.name=="Bibliography")
ok("References: PDF = Word = 53",nb==53 and nw==53,f"PDF {nb}, Word {nw}")
ok("Word refs: Springer style with DOIs","Gneiting T, Raftery AE (2007)" in dt and "https://doi.org/10.5281/zenodo.7970649" in dt)
ok("Unresolved markers: none","??" not in pdf and "[@" not in dt and "\\cite" not in dt and "::: REFS" not in dt)
ok("Figures/: 10 EPS + 10 PDF",len(glob.glob(S+"Figures/Fig*.eps"))==10 and len(glob.glob(S+"Figures/Fig*.pdf"))==10)
ok("Figures/ PDFs = manuscript figures",all(md5(f"serra/figures_serra/Fig{i}.pdf")==md5(f"{S}Figures/Fig{i}.pdf") for i in range(1,11)))
ok("Package PDF = compiled manuscript",md5(S+"Manuscript.pdf")==md5("serra/Manuscript.pdf"))
ok("Package Word = built Word",md5(S+"Manuscript.docx")==md5("serra/word/Manuscript_SERRA.docx"))
ok("Table audit: 0 issues","TABLE ISSUES: 0" in sh("python3 audit_tables_serra.py"))
t=sh("python3 audit_text_serra.py"); ok("Text audit: 24 + 18 checks, 0 issues","CHECKS OK: 24   ISSUES: 0" in t and "ADDITIONS OK: 18   ISSUES: 0" in t)
z=zipfile.ZipFile("SERRA_submission.zip"); ok("Zip contains latest PDF",hashlib.md5(z.read("SERRA_submission/Manuscript.pdf")).hexdigest()==md5(S+"Manuscript.pdf"))
ok("Zip contains latest Word",hashlib.md5(z.read("SERRA_submission/Manuscript.docx")).hexdigest()==md5(S+"Manuscript.docx"))
rz=zipfile.ZipFile("B8_repo_update_v1.3.0.zip").namelist()
ok("Repo v1.3.0 has new scripts/results/NWP data",all(any(n.endswith(x) for n in rz) for x in ["run/run_nwp.py","run/run_refit.py","run/run_linqr.py","results/nwp_summary.json","results/refit_summary.json","results/consistency.json","results/nongauss.json","data/nwp/README.md","results/PROVENANCE.md"]))

ok("ORCID on title page (PDF)","0009-0004-5659-2110" in pn); ok("ORCID on title page (Word)","0009-0004-5659-2110" in dn)
ok("Journal names abbreviated (Word refs)","J Am Stat Assoc 102" in dt and "Journal of the American Statistical Association" not in dt)
ok("Journal names abbreviated (PDF refs)","J Am Stat Assoc" in pn and "Journal of the American Statistical Association" not in pn)
import fitz
low=[]
for i in range(1,11):
    p=fitz.open(f"SERRA_submission/Figures/Fig{i}.pdf")[0]; w=p.rect.width/72; s_=((84 if i in (2,9) else 174)/25.4)/w; reg=[]
    for b in p.get_text("dict")["blocks"]:
        for l in b.get("lines",[]):
            prev=None
            for sp in l["spans"]:
                if not sp["text"].strip(): continue
                if prev and sp["size"]<0.8*prev: continue
                reg.append(sp["size"]); prev=sp["size"]
    if min(reg)*s_<7.95: low.append(f"Fig{i} {min(reg)*s_:.1f}")
ok("All figure lettering >= 8 pt at print width",not low,", ".join(low))
from PIL import Image
Image.MAX_IMAGE_PIXELS=None

_tx=open("serra/Manuscript.tex").read(); body=re.sub(r"\s+"," ",_tx[_tx.index("\\section{Introduction}"):_tx.index("\\bibliographystyle")])
for ab,full in [("PV","photovoltaic (PV)"),("CAMELS","Catchment Attributes and Meteorology"),("DKASC","Desert Knowledge Australia Solar Centre (DKASC)"),
                ("HKUST","Technology (HKUST)"),("ECMWF","Medium-Range Weather Forecasts (ECMWF)"),
                ("GFS","Global Forecast System (GFS)"),("NOAA","Administration (NOAA)"),("RAWRIS","System (RAWRIS)"),("SCADA","(SCADA)"),("CI","confidence intervals (CIs)")]:
    i1=body.find(full); i2=re.search(r"\b"+ab+r"\b",body)
    ok(f"Abbreviation defined at first use: {ab}", i1>=0 and (i2 is None or i1<=i2.start()+len(full)))
ok("Statements: ethics approval","Ethics approval" in pn.split("Statements and Declarations")[-1])
ok("Statements: consent","Consent to participate" in pn.split("Statements and Declarations")[-1])
ok("Main-text figures numbered consecutively (no A1 figure numbering)","Fig. A" not in pn)
mx=[];fn=set();lw=[];hh=[]
for i in range(1,11):
    p=fitz.open(f"SERRA_submission/Figures/Fig{i}.pdf")[0]; w=p.rect.width/72; s_=((84 if i in (2,9) else 174)/25.4)/w
    for b in p.get_text("dict")["blocks"]:
        for l in b.get("lines",[]):
            for sp in l["spans"]:
                if sp["text"].strip(): fn.add(sp["font"].split("+")[-1].split("-")[0]); mx.append(sp["size"]*s_)
    lw+= [d["width"]*s_ for d in p.get_drawings() if d.get("width")]; hh.append(p.rect.height/72*25.4*s_)
ok("Figure lettering <= 12 pt at print width",max(mx)<=12.05,f"max {max(mx):.1f}")
ok("Figure font Arial-metric (Liberation Sans)",{f for f in fn if not f.startswith("Cmsy")}=={"LiberationSans"},str(sorted(fn)))
ok("Figure lines >= 0.3 pt (0.1 mm)",min(lw)>=0.3,f"min {min(lw):.2f}")
ok("Figure height <= 234 mm",max(hh)<=234,f"max {max(hh):.0f} mm")
ok("TIFF 1200 dpi at print width",all(abs(Image.open(f"B8_SERRA_figures/TIFF_1200dpi/Fig{i}.tiff").width/((84 if i in (2,9) else 174)/25.4)-1200)<5 for i in range(1,11)))
for s,n,d in R: print(f"{s} {n}"+(f"  [{d}]" if d else ""))
print(f"\nPASSED {sum(r[0]=='✅' for r in R)}/{len(R)}")
