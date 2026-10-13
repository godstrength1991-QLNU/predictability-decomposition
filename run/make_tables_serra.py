import re, csv, os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
tex=open("serra/Manuscript.tex").read()
tex=re.sub(r"\\input\{(tab_[a-z]+)\}",lambda m:open("serra/"+m.group(1)+".tex").read(),tex)
aux=open("serra/Manuscript.aux").read(); lab={m.group(1):m.group(2) for m in re.finditer(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}",aux)}
tex=re.sub(r"\\eqref\{([^}]+)\}",lambda m:"("+lab[m.group(1)]+")",tex); tex=re.sub(r"~?\\ref\{([^}]+)\}",lambda m:" "+lab[m.group(1)],tex)
NAME={"tab:notation":"notation","tab:main":"decomposition_by_domain","tab:redundancy":"redundancy","tab:pe":"permutation_entropy",
      "tab:uncertainty":"uncertainty_decomposition","tab:estimators":"two_estimator_families","tab:rolling":"rolling_origin",
      "tab:drift":"drift_correction","tab:robust":"spatial_sample_robustness","tab:capacity":"capacity_robustness","tab:raw":"raw_CRPS"}
U={r"\rendo":"ρ_endo",r"\rexo":"ρ_exo",r"\Cclim":"CRPS_clim",r"\Cpers":"CRPS_persist",r"\Chist":"CRPS_hist",r"\Cfull":"CRPS_full",
   r"\rho":"ρ",r"\sigma":"σ",r"\mid":"|",r"\ast":"*",r"\cdot":"·",r"\approx":"≈",r"\geq":"≥",r"\tilde R":"R̃",r"\mathrm":"",r"\mathsf":"",r"\text":"",r"\textsf":""}
def clean(c):
    c=c.replace("$-$","−").replace("\\%","%").replace("\\,"," ").replace("~"," ").replace("\\_","_")
    c=re.sub(r"\\mathrm\{sd\}_\{\\mathrm\{([a-z]+)\}\}\^2",r"sd_\1²",c)
    for k,v in sorted(U.items(),key=lambda x:-len(x[0])): c=c.replace(k,v)
    c=re.sub(r"\^\{-1\}","⁻¹",c); c=re.sub(r"\^\{1/2\}","^(1/2)",c); c=re.sub(r"\\textbf\{([^}]*)\}",r"\1",c)
    c=c.replace("$","").replace("{","").replace("}","").replace("\\","")
    c=re.sub(r"(-)(\d)",r"−\2",c) if re.fullmatch(r"\s*-\d.*",c) else c
    return re.sub(r"\s+"," ",c).strip()
def cells(row):
    out=[]
    for c in row.split("&"):
        m=re.match(r"\s*\\multicolumn\{(\d+)\}\{[^}]*\}\{(.*)\}\s*$",c.strip(),re.S)
        if m: out+= [clean(m.group(2))]+[""]*(int(m.group(1))-1)
        else: out.append(clean(c))
    return out
SUB=re.compile(r"(ρ|CRPS|R̃|p|sd)_(endo|exo|clim|persist|hist|full|shr|•|Sh|refit|boot|tot)")
def write(par,text,size=10,bold=False):
    pos=0
    for m in SUB.finditer(text):
        r=par.add_run(text[pos:m.start()]+m.group(1)); r.font.size=Pt(size); r.font.name="Times New Roman"; r.bold=bold
        s=par.add_run(m.group(2)); s.font.subscript=True; s.font.size=Pt(size); s.font.name="Times New Roman"; s.bold=bold; pos=m.end()
    r=par.add_run(text[pos:]); r.font.size=Pt(size); r.font.name="Times New Roman"; r.bold=bold
def border(parent,tag,sides):
    b=parent.find(qn(tag))
    if b is not None: parent.remove(b)
    b=OxmlElement(tag)
    for side,val in sides:
        e=OxmlElement("w:"+side); e.set(qn("w:val"),val)
        if val!="nil": e.set(qn("w:sz"),"8"); e.set(qn("w:space"),"0"); e.set(qn("w:color"),"000000")
        b.append(e)
    parent.append(b)
done=[]
for t in re.findall(r"\\begin\{table\}.*?\\end\{table\}",tex,re.S):
    key=re.search(r"\\label\{([^}]+)\}",t).group(1); no=lab[key]
    fname=f"Table{no.replace('.','')}_{NAME[key]}"
    cap=clean(re.search(r"\\caption\{(.*?)\}\s*\\label",t,re.S).group(1))
    tab=re.search(r"\\begin\{tabular\}\{[^}]*\}(.*?)\\end\{tabular\}",t,re.S).group(1)
    tab=re.sub(r"\\(toprule|midrule|bottomrule)|\\cmidrule\([^)]*\)\{[^}]*\}|\\addlinespace(\[[^\]]*\])?","",tab)
    rows=[cells(r) for r in re.split(r"\\\\",tab) if r.strip()]
    ncol=max(len(r) for r in rows); rows=[r+[""]*(ncol-len(r)) for r in rows]
    with open(f"B8_SERRA_tables/CSV/{fname}.csv","w",newline="",encoding="utf-8-sig") as f:
        w=csv.writer(f); w.writerow([f"Table {no} {cap}"]); w.writerows(rows)
    body=re.sub(r"\\label\{[^}]+\}","",t).replace("\\caption{",f"\\caption{{Table {no} ",1)
    open(f"B8_SERRA_tables/LaTeX/{fname}.tex","w").write(body.replace("[t]","[h]")+"\n")
    d=Document(); sec=d.sections[0]
    if ncol>=6: sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width,sec.page_height=sec.page_height,sec.page_width
    p=d.add_paragraph(); write(p,f"Table {no}",bold=True); write(p," "+cap)
    T=d.add_table(rows=len(rows),cols=ncol); tp=T._tbl.tblPr
    ww=OxmlElement("w:tblW"); ww.set(qn("w:type"),"pct"); ww.set(qn("w:w"),"5000"); tp.append(ww)
    border(tp,"w:tblBorders",[("top","single"),("left","nil"),("bottom","single"),("right","nil"),("insideH","nil"),("insideV","nil")])
    usable=(sec.page_width-sec.left_margin-sec.right_margin)/914400
    lens=[max(5,max(len(r[j]) for r in rows)) for j in range(ncol)]; lens[0]=int(lens[0]*1.5); tot=sum(lens); T.autofit=False
    grid=T._tbl.tblGrid
    for j,gc in enumerate(grid.findall(qn("w:gridCol"))): gc.set(qn("w:w"),str(int(usable*lens[j]/tot*1440)))
    lay=OxmlElement("w:tblLayout"); lay.set(qn("w:type"),"fixed"); tp.append(lay)
    hdr=1 if any(r[0]=="" for r in rows[:1]) else 0
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            c=T.cell(i,j); c.width=Inches(usable*lens[j]/tot); par=c.paragraphs[0]
            par.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after=Pt(1); write(par,val)
            if i==hdr: border(c._tc.get_or_add_tcPr(),"w:tcBorders",[("bottom","single")])
    if hdr==1:                                   # merge grouped header cells (from \multicolumn)
        j=1
        while j<ncol:
            if rows[0][j]:
                k=j+1
                while k<ncol and rows[0][k]=="": k+=1
                if k-j>1:
                    m=T.cell(0,j).merge(T.cell(0,k-1))
                    for extra in m.paragraphs[1:]: extra._p.getparent().remove(extra._p)
                    m.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
                j=k
            else: j+=1
    d.save(f"B8_SERRA_tables/Word/{fname}.docx"); done.append((no,fname,len(rows),ncol))
for x in done: print(f"Table {x[0]:4s} {x[1]:42s} {x[2]} rows x {x[3]} cols")
